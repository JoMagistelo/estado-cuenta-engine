from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from exporters.excel.batch_exporter import export_batch_excel
from models.datos_cuenta import DatosCuenta
from models.estado_cuenta import EstadoCuenta
from models.ocr_review import OCRCandidate, OCRReview
from models.otros_productos import OtrosProductos
from models.processing_result import ProcessingResult
from models.resumen_financiero import ResumenFinanciero


def _estado_cuenta(movimientos) -> EstadoCuenta:
    estado = EstadoCuenta(
        datos_cuenta=DatosCuenta(
            producto_principal="Cuenta de prueba",
            periodo_inicio="01/01/2026",
            periodo_fin="31/01/2026",
            fecha_corte="31/01/2026",
            numero_cuenta="123456",
            numero_cliente=None,
            clabe=None,
            nombre_cliente="Prueba",
            rfc=None,
        ),
        otros_productos=OtrosProductos(
            contrato=None,
            producto=None,
            tasa_interes_anual=None,
            gat_nominal_anual=None,
            gat_real_anual=None,
            total_comisiones=None,
        ),
        resumen_financiero=ResumenFinanciero(
            saldo_promedio=0.0,
            dias_periodo=31,
            tasa_bruta_anual=0.0,
            saldo_promedio_gravable=0.0,
            intereses_a_favor=0.0,
            isr_retenido=0.0,
            cheques_pagados=0,
            manejo_cuenta=0.0,
            cargos_objetados=0.0,
            abonos_objetados=0.0,
            saldo_anterior=0.0,
            depositos_abonos=0.0,
            retiros_cargos=0.0,
            saldo_final=0.0,
            saldo_promedio_minimo_mensual=0.0,
            saldo_global=0.0,
        ),
        movimientos=[],
    )
    # Python no impide que un parser/reader excepcional entregue None aunque
    # el modelo use una lista por defecto. Cubrimos ambos bordes de entrada.
    estado.movimientos = movimientos
    return estado


def _result(movimientos) -> ProcessingResult:
    return ProcessingResult(
        file_name="ocr_sin_movimientos.pdf",
        bank_key="HSBC",
        estado_cuenta=_estado_cuenta(movimientos),
        raw_text="",
        normalized_text="",
        processing_method="OCR",
        ocr_engine="paddleocr",
    )


@pytest.mark.parametrize("movimientos", [[], None])
def test_excel_export_preserves_result_when_there_are_zero_movements(
    tmp_path: Path,
    movimientos,
) -> None:
    output = export_batch_excel([_result(movimientos)], tmp_path / "sin_movimientos.xlsx")

    workbook = load_workbook(output, data_only=True)
    assert workbook["Movimientos"]["A1"].value == "Sin información"
    assert workbook["Datos de la Cuenta"].max_row == 2
    assert workbook["Resumen Financiero"].max_row == 2


def test_excel_export_accepts_confirmed_dual_ocr_candidate_with_zero_movements(
    tmp_path: Path,
) -> None:
    tesseract = OCRCandidate(
        engine="tesseract",
        estado_cuenta=_estado_cuenta([]),
        document=SimpleNamespace(raw_text="", normalized_text=""),
        validaciones=[],
    )
    paddle = OCRCandidate(
        engine="paddleocr",
        estado_cuenta=_estado_cuenta(None),
        document=SimpleNamespace(raw_text="", normalized_text=""),
        validaciones=[],
    )
    review = OCRReview(
        candidates={"tesseract": tesseract, "paddleocr": paddle},
        recommended_engine="tesseract",
        selected_engine="paddleocr",
        confirmed_engine="paddleocr",
    )
    result = ProcessingResult(
        file_name="ocr_dual_sin_movimientos.pdf",
        bank_key="HSBC",
        estado_cuenta=tesseract.estado_cuenta,
        raw_text="",
        normalized_text="",
        processing_method="OCR",
        ocr_review=review,
        ocr_engine="tesseract",
        ocr_primary_engine="tesseract",
        ocr_secondary_engine="paddleocr",
        fallback_attempted=True,
    )

    output = export_batch_excel([result], tmp_path / "dual_sin_movimientos.xlsx")

    assert result.confirmed_ocr_engine == "paddleocr"
    assert result.ocr_engine == "paddleocr"
    workbook = load_workbook(output, data_only=True)
    assert workbook["Movimientos"]["A1"].value == "Sin información"
    assert workbook["Datos de la Cuenta"].max_row == 2
