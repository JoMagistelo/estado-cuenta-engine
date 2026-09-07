from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from exporters.excel.batch_exporter import export_batch_excel
from models.datos_cuenta import DatosCuenta
from models.estado_cuenta import EstadoCuenta
from models.otros_productos import OtrosProductos
from models.processing_result import ProcessingResult
from models.resumen_financiero import ResumenFinanciero


EXPECTED_ACCOUNT_HEADERS = [
    "id_estado",
    "Nombre del Archivo",
    "Banco",
    "Número de Cliente",
    "Nombre del Cliente",
    "RFC",
    "Producto Principal",
    "Número de Cuenta",
    "CLABE",
    "Periodo de Inicio",
    "Periodo de Fin",
    "Fecha de Corte",
]

EXPECTED_SUMMARY_HEADERS = [
    "id_estado",
    "Nombre del Archivo",
    "Banco",
    "Días del Periodo ",
    "Intereses a Favor",
    "ISR Retenido",
    "Saldo Anterior",
    "Depositos / Abonos (+)",
    "Retiros / Cargos (-)",
    "Saldo Final",
    "Saldo Promedio",
    "Saldo Promedio Mínimo Mensual",
    "Saldo Promedio Gravable",
    "Tasa Bruta Anual",
    "Cheques Pagados",
    "Manejo de Cuenta",
    "Cargos Objetados",
    "Abonos Objetados",
    "Saldo Global",
]


def _result() -> ProcessingResult:
    estado = EstadoCuenta(
        datos_cuenta=DatosCuenta(
            producto_principal="Cuenta de prueba",
            periodo_inicio="01/01/2026",
            periodo_fin="31/01/2026",
            fecha_corte="31/01/2026",
            numero_cuenta="123456",
            numero_cliente="CLIENTE-1",
            clabe="012345678901234567",
            nombre_cliente="Persona de prueba",
            rfc="RFC010101ABC",
        ),
        otros_productos=OtrosProductos(
            contrato=None,
            producto=None,
            tasa_interes_anual=None,
            gat_nominal_anual=None,
            gat_real_anual=None,
            total_comisiones=None,
        ),
        resumen_financiero=ResumenFinanciero(
            saldo_promedio=90.0,
            dias_periodo=31,
            tasa_bruta_anual=1.0,
            saldo_promedio_gravable=80.0,
            intereses_a_favor=2.0,
            isr_retenido=0.5,
            cheques_pagados=0,
            manejo_cuenta=0.0,
            cargos_objetados=0.0,
            abonos_objetados=0.0,
            saldo_anterior=100.0,
            depositos_abonos=50.0,
            retiros_cargos=25.0,
            saldo_final=125.0,
            saldo_promedio_minimo_mensual=70.0,
            saldo_global=125.0,
        ),
        movimientos=[],
    )
    return ProcessingResult(
        file_name="layout.xlsx.pdf",
        bank_key="HSBC",
        estado_cuenta=estado,
        raw_text="",
        normalized_text="",
        processing_method="Digital",
    )


def test_exported_workbook_uses_requested_sheet_order_zoom_and_headers(tmp_path: Path) -> None:
    output = export_batch_excel([_result()], tmp_path / "layout.xlsx")
    workbook = load_workbook(output, data_only=True)

    assert workbook.sheetnames == [
        "Datos de la Cuenta",
        "Resumen Financiero",
        "Movimientos",
        "Otros Productos",
    ]
    assert all(sheet.sheet_view.zoomScale == 55 for sheet in workbook.worksheets)

    account_headers = [cell.value for cell in workbook["Datos de la Cuenta"][1]]
    summary_headers = [cell.value for cell in workbook["Resumen Financiero"][1]]

    assert account_headers == EXPECTED_ACCOUNT_HEADERS
    assert summary_headers == EXPECTED_SUMMARY_HEADERS
