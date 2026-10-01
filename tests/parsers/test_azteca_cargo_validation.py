"""Validación Azteca con estados sintéticos, sin dumps bancarios reales."""

from __future__ import annotations

import ast
import copy
import pickle
from dataclasses import fields, replace
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from engine import ocr_reprocessing, pipeline, statement_processor
from exporters.excel.batch_exporter import export_batch_excel
from models.datos_cuenta import DatosCuenta
from models.estado_cuenta import EstadoCuenta
from models.movimiento import Movimiento
from models.otros_productos import OtrosProductos
from models.processing_result import ProcessingResult
from models.resumen_financiero import ResumenFinanciero
from parsers.azteca.extractors.movimientos import enrich_movement_metadata_from_concepto
from parsers.azteca.validacion import ResumenFinancieroAzteca
from validators.movimiento_validator import validar_movimientos


CASES = {
    "mar": {
        "saldo_anterior": 3.55,
        "depositos": 2600.00,
        "reportado": 2524.00,
        "cargos": (1000.40, 1524.41),
        "saldo_final": 79.55,
    },
    "feb": {
        "saldo_anterior": 0.92,
        "depositos": 3831.63,
        "reportado": 3829.00,
        "cargos": (1800.20, 2029.23),
        "saldo_final": 3.55,
    },
    "abr": {
        "saldo_anterior": 79.55,
        "depositos": 2892.47,
        "reportado": 2904.00,
        "cargos": (1400.00, 1504.00),
        "saldo_final": 68.02,
    },
}


def _incoming(amount: float) -> Movimiento:
    movement = Movimiento(
        "01/03/2026",
        None,
        "TRANSFERENCIA SPEI A SU FAVOR EMISOR: BBVA MEXICO "
        "CUENTA: 012180000000000001 NOM ORIGI: PERSONA EJEMPLO "
        "RASTREO: TEST-IN-0001 REF: 0000001 CONCEPTO: prueba",
        "ABONO",
        0.0,
        amount,
    )
    return enrich_movement_metadata_from_concepto(movement)


def _outgoing(amount: float, *, detailed: bool = False) -> Movimiento:
    concept = "CARGO SINTETICO"
    if detailed:
        concept = (
            "ORDEN DE TRANSFERENCIA SPEI RECEPTOR: BBVA MEXICO "
            "NOM BENEF: PERSONA EJEMPLO DATO NO VERIFICADO POR ESTA INSTITUCION. "
            "RASTREO: TEST-OUT-0001 REF: 0000002 CONCEPTO: prueba"
        )
    movement = Movimiento("02/03/2026", None, concept, "CARGO", amount, 0.0)
    return enrich_movement_metadata_from_concepto(movement) if detailed else movement


def state(case: str = "mar") -> EstadoCuenta:
    values = CASES[case]
    movements = [
        _incoming(values["depositos"]),
        _outgoing(values["cargos"][0]),
        _outgoing(values["cargos"][1], detailed=True),
    ]
    summary = ResumenFinancieroAzteca(
        saldo_promedio=0.0,
        dias_periodo=31,
        tasa_bruta_anual=0.0,
        saldo_promedio_gravable=0.0,
        intereses_a_favor=0.0,
        isr_retenido=0.0,
        cheques_pagados=0,
        manejo_cuenta=0.0,
        cargos_objetados=0.0,
        abonos_objetados=0.0,
        saldo_anterior=values["saldo_anterior"],
        depositos_abonos=values["depositos"],
        retiros_cargos=values["reportado"],
        saldo_final=values["saldo_final"],
        saldo_promedio_minimo_mensual=0.0,
        saldo_global=values["saldo_final"],
        total_retiros_tabla=values["reportado"],
    )
    return EstadoCuenta(
        datos_cuenta=DatosCuenta(
            producto_principal="GUARDADITO DIGITAL",
            periodo_inicio="01/03/2026",
            periodo_fin="31/03/2026",
            fecha_corte="31/03/2026",
            numero_cuenta="90000000000000",
            numero_cliente="90000000",
            clabe="127180000000000001",
            nombre_cliente="PERSONA PRUEBA",
            rfc="XAXX010101000",
        ),
        otros_productos=OtrosProductos(None, "N/A", None, None, None, None),
        resumen_financiero=summary,
        movimientos=movements,
    )


def cargos_validation(statement: EstadoCuenta):
    return next(
        v
        for v in validar_movimientos(statement.movimientos, statement.resumen_financiero)
        if v.nombre == "Total retiros / cargos"
    )


@pytest.mark.parametrize(
    "case,reported,detail,difference",
    [
        ("mar", 2524.0, 2524.81, 0.81),
        ("feb", 3829.0, 3829.43, 0.43),
    ],
)
def test_known_centavo_omission_is_conditional_and_preserves_both_amounts(
    case,
    reported,
    detail,
    difference,
) -> None:
    statement = state(case)
    before = copy.deepcopy(statement)
    validation = cargos_validation(statement)
    assert validation.correcto and validation.advertencia
    assert validation.esperado == reported
    assert validation.obtenido == pytest.approx(detail)
    assert validation.diferencia == pytest.approx(difference)
    assert "posible omisión de centavos" in validation.mensaje
    assert f"${reported:,.2f}" in validation.mensaje
    assert f"${detail:,.2f}" in validation.mensaje
    assert statement.resumen_financiero.retiros_cargos == reported
    assert statement.resumen_financiero.total_retiros_tabla == reported
    assert statement == before
    assert all(
        v.correcto for v in validar_movimientos(statement.movimientos, statement.resumen_financiero)
    )


def test_exact_synthetic_layout_keeps_normal_validation() -> None:
    statement = state("abr")
    validation = cargos_validation(statement)
    assert validation.correcto and not validation.advertencia
    assert all(
        v.correcto for v in validar_movimientos(statement.movimientos, statement.resumen_financiero)
    )


@pytest.mark.parametrize("case", ["mar", "feb"])
def test_same_figures_in_generic_bank_summary_are_still_a_failure(case) -> None:
    statement = state(case)
    specific = statement.resumen_financiero
    statement.resumen_financiero = ResumenFinanciero(
        **{f.name: getattr(specific, f.name) for f in fields(ResumenFinanciero)}
    )
    validation = cargos_validation(statement)
    assert not validation.correcto and not validation.advertencia


@pytest.mark.parametrize(
    "changes",
    [
        {"total_retiros_tabla": None},
        {"total_retiros_tabla": 2525.0},
        {"retiros_cargos": 2524.01, "total_retiros_tabla": 2524.01, "saldo_final": 79.54},
        {"retiros_cargos": 2523.0, "total_retiros_tabla": 2523.0, "saldo_final": 80.55},
        {"retiros_cargos": 2525.0, "total_retiros_tabla": 2525.0, "saldo_final": 78.55},
        {"depositos_abonos": 2600.01, "saldo_final": 79.56},
        {"saldo_final": 79.56},
        {"saldo_anterior": None},
        {"saldo_final": None},
        {"total_retiros_tabla": float("nan")},
        {"saldo_final": float("inf")},
    ],
)
def test_missing_or_conflicting_evidence_keeps_the_failure(changes) -> None:
    statement = state()
    statement.resumen_financiero = replace(statement.resumen_financiero, **changes)
    validation = cargos_validation(statement)
    assert not validation.correcto and not validation.advertencia


@pytest.mark.parametrize(
    "changes",
    [
        {"fecha_operacion": ""},
        {"concepto": ""},
        {"cargo": 0.0},
        {"abono": 10.0},
        {"cargo": 3.001},
    ],
)
def test_partial_or_conflicting_movements_do_not_waive_the_difference(changes) -> None:
    statement = state()
    statement.movimientos[2] = replace(statement.movimientos[2], **changes)
    validation = cargos_validation(statement)
    assert not validation.correcto and not validation.advertencia


def test_unexplained_whole_peso_difference_stays_a_failure() -> None:
    statement = state()
    statement.movimientos.append(
        Movimiento("28/03/2026", None, "CARGO ADICIONAL", "CARGO", 1.0, 0.0)
    )
    validation = cargos_validation(statement)
    assert not validation.correcto and not validation.advertencia


@pytest.mark.parametrize(
    "validate",
    [
        statement_processor._validation_results,
        ocr_reprocessing._validations,
        lambda s: pipeline._result_validations(s, None),
    ],
)
def test_pipeline_review_and_reprocessing_share_the_same_azteca_policy(validate) -> None:
    statement = state()
    result = next(v for v in validate(statement) if v.nombre == "Total retiros / cargos")
    assert result.correcto and result.advertencia
    assert result.diferencia == pytest.approx(0.81)


def test_centavo_evidence_survives_parallel_process_serialization() -> None:
    statement = pickle.loads(pickle.dumps(state()))
    assert cargos_validation(statement).advertencia
    assert statement.resumen_financiero.retiros_cargos == 2524.0


def test_full_synthetic_concept_and_spei_fields_are_preserved() -> None:
    statement = state("mar")
    assert statement.movimientos[0].concepto == (
        "TRANSFERENCIA SPEI A SU FAVOR EMISOR: BBVA MEXICO "
        "CUENTA: 012180000000000001 NOM ORIGI: PERSONA EJEMPLO "
        "RASTREO: TEST-IN-0001 REF: 0000001 CONCEPTO: prueba"
    )
    sent = next(m for m in statement.movimientos if m.cargo and m.clave_rastreo)
    assert sent.concepto == (
        "ORDEN DE TRANSFERENCIA SPEI RECEPTOR: BBVA MEXICO "
        "NOM BENEF: PERSONA EJEMPLO DATO NO VERIFICADO POR ESTA INSTITUCION. "
        "RASTREO: TEST-OUT-0001 REF: 0000002 CONCEPTO: prueba"
    )
    assert statement.movimientos[0].concepto_original == "prueba"
    assert sent.concepto_original == "prueba"


def test_excel_keeps_synthetic_concepts_and_original_financial_values(tmp_path) -> None:
    statement = state()
    result = ProcessingResult("Azteca.pdf", "azteca", statement, "", "", processing_method="OCR")
    path = export_batch_excel([result], tmp_path / "azteca.xlsx")
    workbook = load_workbook(path)
    movement_sheet = workbook["Movimientos"]
    headers = [c.value for c in movement_sheet[1]]
    for index, movement in enumerate(statement.movimientos, start=2):
        assert movement_sheet.cell(index, headers.index("Concepto") + 1).value == movement.concepto
        assert (
            movement_sheet.cell(index, headers.index("Concepto Original") + 1).value
            == movement.concepto_original
        )
        assert movement_sheet.cell(index, headers.index("Cargo") + 1).value == movement.cargo
    summary = workbook["Resumen Financiero"]
    columns = [c.value for c in summary[1]]
    assert summary.cell(2, columns.index("Retiros / Cargos (-)") + 1).value == 2524.0
    workbook.close()


@pytest.mark.parametrize("case,expected_icon", [("mar", "⚠️"), ("abr", "✅")])
def test_flet_validation_card_shows_warning_or_exact_conciliation(case, expected_icon) -> None:
    root = Path(__file__).parents[2]
    module = ast.parse((root / "app" / "main_flet.py").read_text(encoding="utf-8"))
    card = next(
        n
        for n in ast.walk(module)
        if isinstance(n, ast.FunctionDef) and n.name == "validation_card"
    )

    def control(name):
        return lambda *args, **kwargs: {"control": name, "args": args, "kwargs": kwargs}

    ft = SimpleNamespace(**{name: control(name) for name in ["Container", "Row", "Column", "Text"]})
    ft.Colors = SimpleNamespace(
        ON_SURFACE_VARIANT="grey",
        ORANGE="orange",
        GREEN="green",
        RED="red",
        OUTLINE_VARIANT="outline",
    )
    ft.FontWeight = SimpleNamespace(BOLD="bold")
    ft.Border = SimpleNamespace(all=lambda *args: None)
    namespace = {
        "ft": ft,
        "validation": lambda result, name: result,
        "format_money": lambda value: f"${value:,.2f}",
    }
    exec(compile(ast.Module(body=[card], type_ignores=[]), "validation_card", "exec"), namespace)
    validation = cargos_validation(state(case))
    rendered = namespace["validation_card"](validation, validation.nombre, "Validación cargos")
    rendered_row = rendered["args"][0]["args"][0]
    assert rendered_row[0]["args"][0] == expected_icon
    texts = rendered_row[1]["args"][0]
    if validation.advertencia:
        assert texts[0]["kwargs"]["color"] == "orange"
        assert texts[1]["args"][0] == validation.mensaje
        assert "$2,524.00" in texts[1]["args"][0]
        assert "$2,524.81" in texts[1]["args"][0]
    else:
        assert texts[0]["kwargs"]["color"] == "green"
        assert texts[1]["args"][0] == "Conciliación correcta"
